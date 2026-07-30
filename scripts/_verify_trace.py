#!/usr/bin/env python3
"""Verify analysis trace deliverables for regression triage."""

from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl
import pandas as pd


def verify_banh_chung(tid: str) -> dict:
    base = Path(f"/app/data/artifacts/{tid}")
    out: dict = {"trace_id": tid, "kind": "banh_chung", "ok": False}
    if not (base / "out/data_agent_trace.json").exists():
        out["error"] = "missing_trace"
        return out
    trace = json.loads((base / "out/data_agent_trace.json").read_text())
    ws = base / "ws"
    sale = len(pd.read_parquet(ws / "sale_lines.parquet")) if (ws / "sale_lines.parquet").exists() else 0
    comp = (
        len(pd.read_parquet(ws / "sale_lines_all_bill_lines.parquet"))
        if (ws / "sale_lines_all_bill_lines.parquet").exists()
        else 0
    )
    cust = len(pd.read_parquet(ws / "export_customers.parquet")) if (ws / "export_customers.parquet").exists() else 0
    xlsx = base / "out/analysis_result.xlsx"
    sheets: dict[str, int] = {}
    if xlsx.exists():
        wb = openpyxl.load_workbook(xlsx, read_only=True)
        sheets = {s: sum(1 for _ in wb[s].iter_rows(min_row=2)) for s in wb.sheetnames}
    out.update(
        {
            "sale_lines": sale,
            "companion_lines": comp,
            "customers": cust,
            "xlsx_sheets": sheets,
            "finalize": (trace.get("finalize") or {}).get("action"),
        }
    )
    out["ok"] = (
        sale >= 10
        and comp >= 50
        and cust >= 10
        and sheets.get("bills", 0) >= 50
        and sheets.get("customers", 0) >= 10
    )
    return out


def verify_top10(tid: str) -> dict:
    base = Path(f"/app/data/artifacts/{tid}")
    out: dict = {"trace_id": tid, "kind": "top10_revenue", "ok": False}
    if not (base / "out/data_agent_trace.json").exists():
        out["error"] = "missing_trace"
        return out
    xlsx = base / "out/analysis_result.xlsx"
    if not xlsx.exists():
        out["error"] = "missing_xlsx"
        return out
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    best_rows = 0
    best_sheet = None
    for s in wb.sheetnames:
        n = sum(1 for _ in wb[s].iter_rows(min_row=2))
        if n > best_rows:
            best_rows = n
            best_sheet = s
    out["best_sheet"] = best_sheet
    out["data_rows"] = best_rows
    out["sheets"] = list(wb.sheetnames)
    # Expect ~10 ranked SKU rows with revenue/qty columns
    out["ok"] = 5 <= best_rows <= 15
    return out


def main() -> int:
    if len(sys.argv) < 3:
        print("usage: _verify_trace.py <banh_chung|top10> <trace_id>", file=sys.stderr)
        return 2
    kind, tid = sys.argv[1], sys.argv[2]
    result = verify_banh_chung(tid) if kind == "banh_chung" else verify_top10(tid)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result.get("ok") else 1


if __name__ == "__main__":
    raise SystemExit(main())
