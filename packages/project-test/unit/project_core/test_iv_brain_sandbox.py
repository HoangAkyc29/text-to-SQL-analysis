"""Unit tests for Agent IV op-based brain (no sandbox scripts)."""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
import pytest

from project_core.domain.contracts.workflow import PermissionsSnapshot


def _perms() -> PermissionsSnapshot:
    return PermissionsSnapshot(
        actor_id="user-1",
        role="hq_analyst",
        allowed_tables=["STRANS"],
        tool_grants=["tool:*"],
    )


def test_iv_brain_run_op_exports(tmp_path, monkeypatch):
    from project_core.domain.access.context_policy import ContextPolicy
    from project_core.domain.analysis import iv_brain
    from project_core.domain.contracts.brief import AnalysisBrief

    df = pd.DataFrame({"sku": ["a", "b"], "qty": [1, 2]})
    raw = tmp_path / "raw"
    raw.mkdir()
    path = raw / "q0.parquet"
    df.to_parquet(path, index=False)
    out = tmp_path / "out"
    out.mkdir()

    class FakePlanner:
        def __init__(self, *a, **k):
            self.tokens = 2
            self._n = 0

        def next_decision(self, state):
            self._n += 1
            if self._n == 1:
                return {
                    "decision": "run_op",
                    "thought": "export",
                    "op": {
                        "op_id": "export_csv",
                        "dataset": "q0",
                        "args": {"filename": "summary.csv"},
                    },
                }
            return {
                "decision": "finalize",
                "status": "complete",
                "insight_vi": "ok",
                "headline_metrics": {"qty": 3},
            }

    monkeypatch.setattr(iv_brain, "AnalysisPlanner", FakePlanner)

    payload = iv_brain.run_analysis_brain(
        brief=AnalysisBrief(intent="test", metrics=["qty"], output_format=["table"]),
        manifest={"queries": [{"path": str(path), "row_count": 2}]},
        profile={"row_count": 2},
        out_dir=str(out),
        max_steps=4,
        query_meta=[{"role": "main"}],
        recipe_candidates=[],
        domain_rules_excerpt="",
        permissions=_perms(),
        context_policy=ContextPolicy(),
        llm=object(),
        profile_name="analyst",
        system_prompt="x",
    )
    assert payload["action"] == "complete", payload
    assert payload["artifact_paths"]
    assert Path(payload["artifact_paths"][0]).exists()


def test_iv_brain_rejects_script(tmp_path, monkeypatch):
    from project_core.domain.access.context_policy import ContextPolicy
    from project_core.domain.analysis import iv_brain
    from project_core.domain.contracts.brief import AnalysisBrief

    df = pd.DataFrame({"a": [1]})
    path = tmp_path / "q.parquet"
    df.to_parquet(path, index=False)
    out = tmp_path / "out"
    out.mkdir()

    class FakePlanner:
        def __init__(self, *a, **k):
            self.tokens = 1
            self._n = 0

        def next_decision(self, state):
            self._n += 1
            if self._n == 1:
                return {
                    "decision": "run_step",
                    "step": {"kind": "script", "script": "x=1"},
                }
            return {
                "decision": "finalize",
                "status": "partial",
                "insight_vi": "done",
            }

    monkeypatch.setattr(iv_brain, "AnalysisPlanner", FakePlanner)
    payload = iv_brain.run_analysis_brain(
        brief=AnalysisBrief(intent="test", metrics=["a"]),
        manifest={"queries": [{"path": str(path), "row_count": 1}]},
        profile={"row_count": 1},
        out_dir=str(out),
        max_steps=3,
        query_meta=[{"role": "main"}],
        recipe_candidates=[],
        domain_rules_excerpt="",
        permissions=_perms(),
        context_policy=ContextPolicy(),
        llm=object(),
        profile_name="analyst",
        system_prompt="x",
    )
    # Script rejected; finalize may auto-export CSV so action can be partial.
    assert any(
        "script_ops_disabled" in str(c)
        for c in (payload.get("caveats") or [])
    ) or any(
        "script_ops_disabled" in str(s.get("error") or "")
        for s in (payload.get("steps_trace") or [])
    )
    assert payload["action"] in {"data_feedback", "partial", "complete"}
    if payload["action"] == "data_feedback":
        assert payload.get("data_feedback")
    else:
        assert payload.get("artifact_paths"), "auto-export should produce artifacts after script reject"


def test_analysis_planner_logs_preview_on_invalid_json(caplog):
    from project_core.domain.analysis.iv_brain import AnalysisPlanner

    class FakeResult:
        content = ""
        usage_tokens = 0

    class FakeLLM:
        def chat(self, **kwargs):
            return FakeResult()

    planner = AnalysisPlanner(FakeLLM(), "analyst", "system")
    with caplog.at_level("WARNING"):
        with pytest.raises(json.JSONDecodeError):
            planner.next_decision({"x": 1})
    assert "invalid JSON content preview" in caplog.text


def test_iv_reasoning_state_seeds_checklist_and_invalidates_verification():
    from project_core.domain.contracts.brief import AnalysisBrief
    from project_core.domain.contracts.iv_reasoning import (
        IVReasoningPhase,
        IVReasoningState,
        IVVerificationStatus,
    )

    state = IVReasoningState.from_brief(
        AnalysisBrief(
            intent="compare",
            metrics=["qty"],
            dimensions=["store"],
            output_format=["chart"],
        )
    )
    assert {item.item_id for item in state.checklist} == {
        "intent",
        "metric:0",
        "dimension:0",
        "format:0",
    }
    state.advance_to(IVReasoningPhase.VERIFY)
    state.record_verification(
        passed=True,
        artifact_checks={"out.csv": True},
        coverage_gaps=[],
        checked_items=["intent"],
    )
    state.record_mutation("filter_rows")
    assert state.phase == IVReasoningPhase.EXECUTE
    assert state.revision.number == 1
    assert state.verification.status == IVVerificationStatus.UNVERIFIED


def test_iv_brain_direct_finalize_requires_coverage(tmp_path, monkeypatch):
    from project_core.domain.access.context_policy import ContextPolicy
    from project_core.domain.analysis import iv_brain
    from project_core.domain.contracts.brief import AnalysisBrief

    path = tmp_path / "q.parquet"
    pd.DataFrame({"other": [1]}).to_parquet(path, index=False)
    out = tmp_path / "out"
    out.mkdir()

    class FakePlanner:
        def __init__(self, *args, **kwargs):
            self.tokens = 0

        def next_decision(self, state):
            return {"decision": "finalize", "status": "partial"}

    monkeypatch.setattr(iv_brain, "AnalysisPlanner", FakePlanner)
    payload = iv_brain.run_analysis_brain(
        brief=AnalysisBrief(intent="test", metrics=["requested_metric"]),
        manifest={"queries": [{"path": str(path), "row_count": 1}]},
        profile={"row_count": 1},
        out_dir=str(out),
        max_steps=1,
        query_meta=[{"role": "main"}],
        recipe_candidates=[],
        domain_rules_excerpt="",
        permissions=_perms(),
        context_policy=ContextPolicy(),
        llm=object(),
        profile_name="analyst",
        system_prompt="x",
    )
    assert payload["action"] == "data_feedback"
    assert "missing_metric:requested_metric" in payload["reasoning_state"]["verification"][
        "coverage_gaps"
    ]
    assert payload["headline_metrics"]["row_count"] == 1
    assert payload["artifact_manifests"]
    assert payload["artifact_manifests"][0]["validation_status"] == "valid"


def test_iv_brain_aligns_executed_query_meta_by_query_index(tmp_path, monkeypatch):
    from project_core.domain.access.context_policy import ContextPolicy
    from project_core.domain.analysis import iv_brain
    from project_core.domain.contracts.brief import AnalysisBrief, BriefRequirement

    path = tmp_path / "main-query.parquet"
    pd.DataFrame({"quantity": [2, 4]}).to_parquet(path, index=False)
    out = tmp_path / "out-aligned"
    out.mkdir()

    class FakePlanner:
        def __init__(self, *args, **kwargs):
            self.tokens = 0

        def next_decision(self, state):
            return {"decision": "finalize", "status": "partial"}

    monkeypatch.setattr(iv_brain, "AnalysisPlanner", FakePlanner)
    payload = iv_brain.run_analysis_brain(
        brief=AnalysisBrief(
            intent="show quantity",
            metrics=["quantity"],
            output_format=["table"],
            requirements=[
                BriefRequirement(
                    requirement_id="metric:0",
                    kind="metric",
                    key="quantity",
                    source="explicit",
                    required=True,
                    evidence_quote="quantity",
                    value="quantity",
                )
            ],
        ),
        manifest={
            "queries": [
                {
                    "path": str(path),
                    "query_index": 1,
                    "row_count": 2,
                }
            ]
        },
        profile={"row_count": 2},
        out_dir=str(out),
        max_steps=1,
        query_meta=[
            {"role": "probe", "purpose": "resolve"},
            {
                "role": "main",
                "purpose": "quantity",
            },
        ],
        recipe_candidates=[],
        domain_rules_excerpt="",
        permissions=_perms(),
        context_policy=ContextPolicy(),
        llm=object(),
        profile_name="analyst",
        system_prompt="x",
    )
    assert payload["action"] == "complete"
    assert all(
        "unmapped_requirement:metric:0" not in caveat
        for caveat in payload.get("caveats", [])
    )


def test_iv_brain_planner_turn_budget_is_separate_from_op_count(tmp_path, monkeypatch):
    from project_core.domain.access.context_policy import ContextPolicy
    from project_core.domain.analysis import iv_brain
    from project_core.domain.contracts.brief import AnalysisBrief

    path = tmp_path / "q.parquet"
    pd.DataFrame({"qty": [1, 2]}).to_parquet(path, index=False)
    out = tmp_path / "out"
    out.mkdir()

    class FakePlanner:
        def __init__(self, *args, **kwargs):
            self.tokens = 0
            self.calls = 0

        def next_decision(self, state):
            self.calls += 1
            if self.calls == 1:
                return {
                    "decision": "run_op",
                    "op": {
                        "kind": "recipe",
                        "steps": [
                            {
                                "op_id": "select_columns",
                                "dataset": "q0",
                                "save_as": "selected",
                                "args": {"columns": ["qty"]},
                            },
                            {
                                "op_id": "export_csv",
                                "dataset": "selected",
                                "args": {"filename": "summary.csv", "primary": True},
                            },
                        ],
                    },
                }
            return {"decision": "finalize", "status": "complete"}

    monkeypatch.setattr(iv_brain, "AnalysisPlanner", FakePlanner)
    payload = iv_brain.run_analysis_brain(
        brief=AnalysisBrief(intent="test", metrics=["qty"], output_format=["table"]),
        manifest={"queries": [{"path": str(path), "row_count": 2}]},
        profile={"row_count": 2},
        out_dir=str(out),
        max_steps=2,
        query_meta=[{"role": "main"}],
        recipe_candidates=[],
        domain_rules_excerpt="",
        permissions=_perms(),
        context_policy=ContextPolicy(),
        llm=object(),
        profile_name="analyst",
        system_prompt="x",
    )
    assert payload["action"] == "complete"
    assert payload["planner_turns"] == 2
    # Two planner ops plus coverage and export-reload verification checks.
    assert payload["sandbox_steps"] == 4
    assert payload["reasoning_state"]["analysis_ops"] == 2


def test_iv_brain_multi_deliverable_workbook_is_reloaded_and_counted(
    tmp_path, monkeypatch
):
    from project_core.domain.access.context_policy import ContextPolicy
    from project_core.domain.analysis import iv_brain
    from project_core.domain.contracts.brief import AnalysisBrief

    source = tmp_path / "gift.parquet"
    rows = [
        {"sku": f"S{index % 3}", "bill": f"B{index:02}", "qty": 1}
        for index in range(15)
    ]
    pd.DataFrame(rows).to_parquet(source)
    out = tmp_path / "out"
    out.mkdir()

    class FakePlanner:
        def __init__(self, *_args, **_kwargs):
            self.tokens = 0
            self.calls = 0

        def next_decision(self, _state):
            self.calls += 1
            if self.calls == 1:
                return {
                    "decision": "run_op",
                    "op": {
                        "op_id": "groupby_agg",
                        "dataset": "q0",
                        "save_as": "summary",
                        "args": {
                            "by": ["sku"],
                            "aggs": [{"column": "qty", "fn": "sum", "as": "qty"}],
                        },
                    },
                }
            if self.calls == 2:
                return {
                    "decision": "run_op",
                    "op": {
                        "op_id": "export_excel",
                        "args": {
                            "filename": "gift.xlsx",
                            "sheets": {"summary": "summary", "detail": "q0"},
                            "primary": True,
                        },
                    },
                }
            return {"decision": "finalize", "status": "complete"}

    monkeypatch.setattr(iv_brain, "AnalysisPlanner", FakePlanner)
    payload = iv_brain.run_analysis_brain(
        brief=AnalysisBrief(
            intent="summary and detail",
            metrics=["qty"],
            dimensions=["sku"],
            output_format=["excel"],
        ),
        manifest={"queries": [{"path": str(source), "row_count": 15}]},
        profile={"row_count": 15},
        out_dir=str(out),
        max_steps=4,
        max_planner_turns=6,
        query_meta=[{"role": "main"}],
        recipe_candidates=[],
        domain_rules_excerpt="",
        permissions=_perms(),
        context_policy=ContextPolicy(),
        llm=object(),
        profile_name="analyst",
        system_prompt="x",
    )
    assert payload["action"] == "complete", payload
    assert payload["headline_metrics"]["row_count"] == 15
    assert sorted(payload["headline_metrics"]["artifact_rows"].values()) == [3, 15]
    workbook = pd.ExcelFile(payload["artifact_paths"][0])
    assert workbook.sheet_names == ["summary", "detail"]
    assert len(pd.read_excel(workbook, sheet_name="summary")) == 3
    assert len(pd.read_excel(workbook, sheet_name="detail")) == 15


def test_iv_brain_replots_chart_at_most_from_structured_review(tmp_path, monkeypatch):
    from project_core.domain.access.context_policy import ContextPolicy
    from project_core.domain.analysis import iv_brain
    from project_core.domain.contracts.brief import AnalysisBrief
    from project_core.domain.contracts.plot_review import PlotFixArgs, PlotReviewResult

    source = tmp_path / "chart.parquet"
    pd.DataFrame({"label": ["A", "B"], "value": [1, 2]}).to_parquet(source)
    out = tmp_path / "out"
    out.mkdir()

    class FakePlanner:
        def __init__(self, *_args, **_kwargs):
            self.tokens = 0
            self.calls = 0

        def next_decision(self, _state):
            self.calls += 1
            if self.calls == 1:
                return {
                    "decision": "run_op",
                    "op": {
                        "op_id": "plot_chart",
                        "dataset": "q0",
                        "args": {
                            "x": "label",
                            "y": "value",
                            "kind": "bar",
                            "filename": "chart.png",
                        },
                    },
                }
            return {"decision": "finalize", "status": "complete"}

    class FakeReviewer:
        def __init__(self, *_args, **_kwargs):
            self.calls = 0

        def review(self, *_args, **_kwargs):
            self.calls += 1
            if self.calls == 1:
                return PlotReviewResult(
                    verdict="replot",
                    issues=["labels_too_dense"],
                    fix_args=PlotFixArgs(rotate_x_labels=45, width=12, height=6),
                )
            return PlotReviewResult(verdict="pass", issues=[], vision_available=True)

    fake_reviewer = FakeReviewer()
    monkeypatch.setattr(iv_brain, "AnalysisPlanner", FakePlanner)
    monkeypatch.setattr(iv_brain, "ChartReviewer", lambda **_kwargs: fake_reviewer)
    payload = iv_brain.run_analysis_brain(
        brief=AnalysisBrief(
            intent="chart",
            metrics=["value"],
            dimensions=["label"],
            output_format=["chart"],
            chart_spec={"kind": "bar"},
        ),
        manifest={"queries": [{"path": str(source), "row_count": 2}]},
        profile={"row_count": 2},
        out_dir=str(out),
        max_steps=4,
        max_planner_turns=5,
        query_meta=[{"role": "main"}],
        recipe_candidates=[],
        domain_rules_excerpt="",
        permissions=_perms(),
        context_policy=ContextPolicy(),
        llm=object(),
        profile_name="analyst",
        system_prompt="x",
    )
    assert fake_reviewer.calls == 2
    assert payload["action"] == "complete", payload
    assert len(payload["visual_reviews"]) == 2
    assert payload["visual_reviews"][0]["verdict"] == "replot"
    assert payload["visual_reviews"][1]["verdict"] == "pass"
    assert Path(payload["chart_artifacts"][-1]).name == "chart_reviewed_1.png"


def test_iv_brain_accepts_direct_catalog_action_and_registered_bundle(
    tmp_path, monkeypatch
):
    from project_core.domain.access.context_policy import ContextPolicy
    from project_core.domain.analysis import iv_brain
    from project_core.domain.contracts.brief import AnalysisBrief

    source = tmp_path / "direct.parquet"
    pd.DataFrame({"qty": [1, 2]}).to_parquet(source)
    out = tmp_path / "out"
    out.mkdir()

    class FakePlanner:
        def __init__(self, *_args, **_kwargs):
            self.tokens = 0
            self.calls = 0

        def next_decision(self, state):
            self.calls += 1
            if self.calls == 1:
                return {
                    "decision": "cast_column",
                    "dataset": "q0",
                    "columns": ["qty"],
                    "dtype": "float",
                    "save_as": "casted",
                }
            if self.calls == 2:
                return {
                    "decision": "export_csv",
                    "args": {
                        "dataset": "casted",
                        "filename": "direct.csv",
                        "primary": False,
                    },
                }
            if self.calls == 3:
                return {
                    "decision": "bundle_deliverables",
                    "paths": state["artifacts"],
                }
            return {"decision": "finalize", "status": "complete"}

    monkeypatch.setattr(iv_brain, "AnalysisPlanner", FakePlanner)
    payload = iv_brain.run_analysis_brain(
        brief=AnalysisBrief(intent="direct", metrics=["qty"], output_format=["table"]),
        manifest={"queries": [{"path": str(source), "row_count": 2}]},
        profile={"row_count": 2},
        out_dir=str(out),
        max_steps=4,
        max_planner_turns=6,
        query_meta=[{"role": "main"}],
        recipe_candidates=[],
        domain_rules_excerpt="",
        permissions=_perms(),
        context_policy=ContextPolicy(),
        llm=object(),
        profile_name="analyst",
        system_prompt="x",
    )
    assert payload["action"] == "complete", payload
    assert "unknown_decision:bundle_deliverables" not in payload["caveats"]
    assert payload["artifact_manifests"][0]["primary"] is True


def test_iv_brain_stops_repeated_schema_invalid_tool_call_without_spending_op_budget(
    tmp_path,
    monkeypatch,
):
    from project_core.domain.access.context_policy import ContextPolicy
    from project_core.domain.analysis import iv_brain
    from project_core.domain.contracts.brief import AnalysisBrief

    source = tmp_path / "invalid-repeat.parquet"
    pd.DataFrame({"qty": [1, 2]}).to_parquet(source)
    out = tmp_path / "invalid-repeat-out"
    out.mkdir()

    class FakePlanner:
        calls = 0

        def __init__(self, *_args, **_kwargs):
            self.tokens = 0

        def next_decision(self, _state):
            type(self).calls += 1
            return {
                "decision": "cast_column",
                "dataset": "q0",
                "save_as": "never_created",
            }

    monkeypatch.setattr(iv_brain, "AnalysisPlanner", FakePlanner)
    payload = iv_brain.run_analysis_brain(
        brief=AnalysisBrief(intent="inspect", output_format=["table"]),
        manifest={"queries": [{"path": str(source), "row_count": 2}]},
        profile={"row_count": 2},
        out_dir=str(out),
        max_steps=1,
        max_planner_turns=6,
        query_meta=[{"role": "main"}],
        recipe_candidates=[],
        domain_rules_excerpt="",
        permissions=_perms(),
        context_policy=ContextPolicy(),
        llm=object(),
        profile_name="analyst",
        system_prompt="x",
    )
    assert FakePlanner.calls == 2
    assert any(
        step.get("repair_action") == "abandon_schema_invalid_op_and_verify"
        for step in payload["steps_trace"]
    )
    assert payload["reasoning_state"]["analysis_ops"] == 0
    assert payload["verification"]["status"] == "passed"


def test_auto_export_deduplicates_sheets_and_restores_identifier_text(tmp_path):
    from openpyxl import load_workbook

    from project_core.domain.analysis.iv_brain import _ensure_deliverable_exports
    from project_core.domain.analysis.ops.working_set import DatasetWorkingSet
    from project_core.domain.contracts.brief import AnalysisBrief, BriefRequirement

    summary_path = tmp_path / "summary.parquet"
    detail_path = tmp_path / "detail.parquet"
    pd.DataFrame({"ITEM_ID": [9001], "Quantity": ["4.000"]}).to_parquet(summary_path)
    pd.DataFrame(
        {
            "ITEM_ID": [9001, 9001, 9001, 9001],
            "ITEM_CODE": [123, 123, 123, 123],
            "EventDate": pd.to_datetime(
                ["2034-01-04", "2034-01-03", "2034-01-02", "2034-01-01"]
            ),
        }
    ).to_parquet(detail_path)
    working_set = DatasetWorkingSet.from_manifest(
        {
            "queries": [
                {"path": str(summary_path), "ref": "q0", "row_count": 1},
                {"path": str(detail_path), "ref": "q1", "row_count": 4},
            ]
        },
        [
            {"role": "main", "purpose": "summary"},
            {"role": "main", "purpose": "details"},
        ],
        work_dir=tmp_path / "ws-clean",
    )
    working_set.save_frame(
        "q0_copy",
        working_set.get("q0").frame().copy(),
        parents=["q0"],
        op_id="copy",
    )
    brief = AnalysisBrief(
        output_format=["excel"],
        requirements=[
            BriefRequirement(
                requirement_id="filter:0",
                kind="filter",
                key="item_code",
                source="explicit",
                required=True,
                value=["00123"],
            )
        ],
    )
    out = tmp_path / "clean-out"
    out.mkdir()
    exported = _ensure_deliverable_exports(
        working_set,
        brief,
        out_dir=str(out),
        caveats=[],
    )
    assert exported == ["export_excel"]
    workbook = load_workbook(out / "analysis_result.xlsx", data_only=True)
    assert workbook.sheetnames == ["Summary", "Details"]
    assert workbook["Summary"]["B2"].value == "00123"
    assert workbook["Summary"]["B2"].data_type == "s"
    assert workbook["Summary"]["C2"].value == 4
    assert workbook["Summary"]["C2"].data_type == "n"
    assert workbook["Details"]["B2"].value == "00123"
