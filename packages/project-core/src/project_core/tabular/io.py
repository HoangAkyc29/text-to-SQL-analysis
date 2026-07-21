from __future__ import annotations

import csv
import hashlib
from pathlib import Path
from typing import Any, Literal

import pandas as pd

TabularFormat = Literal["parquet", "csv", "xlsx"]


def sha256_file(path: str | Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def detect_format(path: str | Path, declared: str | None = None) -> TabularFormat:
    value = (declared or Path(path).suffix.lstrip(".")).lower()
    if value == "xls":
        raise ValueError("legacy_xls_not_supported")
    if value not in {"parquet", "csv", "xlsx"}:
        raise ValueError(f"unsupported_tabular_format:{value}")
    return value  # type: ignore[return-value]


def read_tabular(
    path: str | Path,
    *,
    format: str | None = None,
    sheet_name: str | int | None = None,
    max_rows: int = 1_000_000,
    max_columns: int = 2_000,
) -> pd.DataFrame:
    source = Path(path)
    fmt = detect_format(source, format)
    if fmt == "parquet":
        df = pd.read_parquet(source)
    elif fmt == "csv":
        # utf-8-sig handles both regular UTF-8 and BOM exports.
        try:
            df = pd.read_csv(source, encoding="utf-8-sig")
        except UnicodeDecodeError:
            df = pd.read_csv(source, encoding="cp1252")
    else:
        df = pd.read_excel(source, sheet_name=sheet_name if sheet_name is not None else 0)
        if isinstance(df, dict):
            raise ValueError("sheet_name_required")
    if len(df) > max_rows:
        raise ValueError(f"row_limit_exceeded:{len(df)}>{max_rows}")
    if len(df.columns) > max_columns:
        raise ValueError(f"column_limit_exceeded:{len(df.columns)}>{max_columns}")
    return df


def inspect_workbook(path: str | Path, *, sample_rows: int = 5) -> dict[str, Any]:
    import openpyxl

    source = Path(path)
    workbook = openpyxl.load_workbook(
        source, read_only=True, data_only=True, keep_links=False
    )
    sheets: list[dict[str, Any]] = []
    try:
        for name in workbook.sheetnames:
            ws = workbook[name]
            rows = ws.iter_rows(values_only=True)
            header = list(next(rows, ()) or ())
            sample: list[list[Any]] = []
            for _, row in zip(range(max(0, min(sample_rows, 20))), rows):
                sample.append(list(row))
            sheets.append(
                {
                    "name": name,
                    "state": ws.sheet_state,
                    "row_count": int(ws.max_row or 0),
                    "column_count": int(ws.max_column or 0),
                    "columns": [str(v) if v is not None else "" for v in header],
                    "sample": sample,
                    "blank_headers": sum(1 for v in header if v in (None, "")),
                    "duplicate_headers": len(header)
                    - len({str(v) for v in header if v not in (None, "")}),
                }
            )
    finally:
        workbook.close()
    return {"path": str(source), "sheet_names": [s["name"] for s in sheets], "sheets": sheets}


def inspect_tabular(
    path: str | Path, *, format: str | None = None, sheet_name: str | None = None
) -> dict[str, Any]:
    source = Path(path)
    fmt = detect_format(source, format)
    if fmt == "xlsx" and sheet_name is None:
        return {
            "format": fmt,
            "byte_size": source.stat().st_size,
            "sha256": sha256_file(source),
            **inspect_workbook(source),
        }
    df = read_tabular(source, format=fmt, sheet_name=sheet_name)
    return {
        "format": fmt,
        "sheet_name": sheet_name,
        "row_count": len(df),
        "columns": [str(c) for c in df.columns],
        "dtypes": {str(c): str(df[c].dtype) for c in df.columns},
        "byte_size": source.stat().st_size,
        "sha256": sha256_file(source),
    }


def neutralize_spreadsheet_formulas(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    for column in out.select_dtypes(include=["object", "string"]).columns:
        out[column] = out[column].map(
            lambda value: "'" + value
            if isinstance(value, str) and value[:1] in {"=", "+", "-", "@"}
            else value
        )
    return out


def sniff_csv_dialect(path: str | Path) -> csv.Dialect:
    with Path(path).open("r", encoding="utf-8-sig", errors="replace") as handle:
        return csv.Sniffer().sniff(handle.read(8192))

